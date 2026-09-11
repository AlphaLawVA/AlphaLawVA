# compare_local_llm_candidate_scores.py
"""
Description: 판례 검색 후보 판례를 로컬 Ollama 모델들이 0~3점으로 평가하게 하고
모델별 점수와 판단 근거를 비교용 CSV 및 엑셀 시트로 저장한다.
Author: choeminju
Date: 2026-09-11
Before:
    - Top10_후보판례_100문항 시트에 후보 판례와 생성요약이 정리된 상태.
    - 로컬 Ollama에 비교할 모델이 설치되어 있고 Ollama API가 실행 가능한 상태.
After:
    - local_data/precedents/evaluation/llm_score_runs/ 아래에 모델 비교 CSV가 생성된다.
    - 평가셋 워크북에 로컬 LLM 점수 비교 시트가 추가 또는 갱신된다.
"""

from __future__ import annotations

import argparse
import csv
import json
import time
import urllib.error
import urllib.request
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK_PATH = (
    PROJECT_ROOT
    / "local_data"
    / "precedents"
    / "evaluation"
    / "AlphaLawVA_precedent_eval_workbook_100_questions.xlsm"
)
DEFAULT_OUTPUT_ROOT = (
    PROJECT_ROOT / "local_data" / "precedents" / "evaluation" / "llm_score_runs"
)
DEFAULT_MODELS = "gemma2:9b,llama3.1:8b"
DEFAULT_SAMPLE_SIZE = 30
DEFAULT_SHEET_NAME = "로컬LLM점수비교_30건"


@dataclass(slots=True)
class CandidateRow:
    """엑셀 후보 판례 시트의 한 행을 비교 실험용 구조로 담는다."""

    source_row: int
    query_id: str
    dispute_type: str
    query: str
    summary: str
    query_specificity: str
    scenario_tags: str
    candidate_rank: str
    precedent_id: str
    case_no: str
    case_name: str
    court: str
    decision_date: str
    matched_models: str
    best_rank: str
    matched_chunk_count: str
    codex_score: str
    codex_reason: str


def parse_args() -> argparse.Namespace:
    """명령행 옵션을 정의한다."""
    parser = argparse.ArgumentParser(
        description="Compare local Ollama models for precedent candidate scoring.",
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK_PATH,
        help="후보 판례와 결과 시트가 들어 있는 XLSM/XLSX 파일 경로.",
    )
    parser.add_argument(
        "--candidate-sheet",
        default="Top10_후보판례_100문항",
        help="후보 판례가 들어 있는 시트 이름.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=None,
        help="결과 CSV를 저장할 폴더. 기본값은 타임스탬프 폴더.",
    )
    parser.add_argument(
        "--result-sheet",
        default=DEFAULT_SHEET_NAME,
        help="워크북에 추가하거나 갱신할 결과 시트 이름.",
    )
    parser.add_argument(
        "--models",
        default=DEFAULT_MODELS,
        help="비교할 Ollama 모델명. 쉼표로 구분한다.",
    )
    parser.add_argument(
        "--sample-size",
        type=int,
        default=DEFAULT_SAMPLE_SIZE,
        help="비교 실험에 사용할 후보 판례 수.",
    )
    parser.add_argument(
        "--ollama-url",
        default="http://127.0.0.1:11434",
        help="Ollama API 기본 URL.",
    )
    parser.add_argument(
        "--timeout",
        type=int,
        default=180,
        help="모델 1회 호출 타임아웃 초.",
    )
    parser.add_argument(
        "--num-ctx",
        type=int,
        default=4096,
        help="Ollama 컨텍스트 길이 옵션.",
    )
    parser.add_argument(
        "--num-predict",
        type=int,
        default=220,
        help="Ollama 응답 최대 토큰 옵션.",
    )
    parser.add_argument(
        "--temperature",
        type=float,
        default=0.0,
        help="평가 일관성을 위한 생성 temperature.",
    )
    parser.add_argument(
        "--prompt-mode",
        choices=["common", "model_specific"],
        default="common",
        help="공통 프롬프트 또는 모델별 보정 프롬프트 사용 여부.",
    )
    parser.add_argument(
        "--skip-workbook",
        action="store_true",
        help="CSV만 만들고 엑셀 시트 갱신은 하지 않는다.",
    )
    return parser.parse_args()


def compact_text(value: Any) -> str:
    """엑셀/CSV 저장을 위해 줄바꿈과 연속 공백을 한 칸으로 정리한다."""
    if value is None:
        return ""
    return " ".join(str(value).split())


def now_utc_iso() -> str:
    """현재 UTC 시각을 ISO 문자열로 반환한다."""
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def make_output_dir(output_dir: Path | None) -> Path:
    """결과 저장 폴더를 생성하고 반환한다."""
    if output_dir is None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_dir = DEFAULT_OUTPUT_ROOT / f"local_llm_score_compare_{timestamp}"
    output_dir.mkdir(parents=True, exist_ok=True)
    return output_dir


def read_candidates(workbook_path: Path, sheet_name: str) -> list[CandidateRow]:
    """후보 판례 시트를 읽어 CandidateRow 목록으로 변환한다."""
    workbook = load_workbook(workbook_path, keep_vba=True, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"시트를 찾지 못했습니다: {sheet_name}")

    sheet = workbook[sheet_name]
    headers = [compact_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    header_index = {name: index for index, name in enumerate(headers)}

    required_headers = [
        "질문ID",
        "분쟁유형",
        "질문",
        "생성요약",
        "질문구체성",
        "시나리오태그",
        "후보순위",
        "판례일련번호",
        "사건번호",
        "사건명",
        "법원",
        "선고일자",
        "검색모델",
        "최고순위",
        "검색청크수",
        "AI점수(0~3)",
        "AI판단라벨근거",
    ]
    missing_headers = [name for name in required_headers if name not in header_index]
    if missing_headers:
        raise ValueError(f"필수 컬럼이 없습니다: {', '.join(missing_headers)}")

    rows: list[CandidateRow] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        query_id = compact_text(row[header_index["질문ID"]])
        precedent_id = compact_text(row[header_index["판례일련번호"]])
        if not query_id or not precedent_id:
            continue
        rows.append(
            CandidateRow(
                source_row=row_number,
                query_id=query_id,
                dispute_type=compact_text(row[header_index["분쟁유형"]]),
                query=compact_text(row[header_index["질문"]]),
                summary=compact_text(row[header_index["생성요약"]]),
                query_specificity=compact_text(row[header_index["질문구체성"]]),
                scenario_tags=compact_text(row[header_index["시나리오태그"]]),
                candidate_rank=compact_text(row[header_index["후보순위"]]),
                precedent_id=precedent_id,
                case_no=compact_text(row[header_index["사건번호"]]),
                case_name=compact_text(row[header_index["사건명"]]),
                court=compact_text(row[header_index["법원"]]),
                decision_date=compact_text(row[header_index["선고일자"]]),
                matched_models=compact_text(row[header_index["검색모델"]]),
                best_rank=compact_text(row[header_index["최고순위"]]),
                matched_chunk_count=compact_text(row[header_index["검색청크수"]]),
                codex_score=compact_text(row[header_index["AI점수(0~3)"]]),
                codex_reason=compact_text(row[header_index["AI판단라벨근거"]]),
            )
        )
    return rows


def select_sample(candidates: list[CandidateRow], sample_size: int) -> list[CandidateRow]:
    """기존 AI점수와 질문구체성이 섞이도록 비교 샘플을 고른다."""
    if sample_size <= 0 or sample_size >= len(candidates):
        return candidates

    groups: dict[tuple[str, str], list[CandidateRow]] = defaultdict(list)
    for row in candidates:
        score = row.codex_score or "blank"
        specificity = row.query_specificity or "unknown"
        groups[(score, specificity)].append(row)

    preferred_order = [
        ("3", "specific_life"),
        ("2", "specific_life"),
        ("1", "specific_life"),
        ("0", "specific_life"),
        ("3", "legal_mixed"),
        ("2", "legal_mixed"),
        ("1", "legal_mixed"),
        ("0", "legal_mixed"),
        ("3", "vague"),
        ("2", "vague"),
        ("1", "vague"),
        ("0", "vague"),
    ]
    selected: list[CandidateRow] = []
    selected_keys: set[tuple[str, str, str]] = set()
    query_counts: Counter[str] = Counter()

    while len(selected) < sample_size:
        progressed = False
        for key in preferred_order:
            for row in groups.get(key, []):
                row_key = (row.query_id, row.precedent_id, row.candidate_rank)
                if row_key in selected_keys or query_counts[row.query_id] >= 2:
                    continue
                selected.append(row)
                selected_keys.add(row_key)
                query_counts[row.query_id] += 1
                progressed = True
                break
            if len(selected) >= sample_size:
                break
        if not progressed:
            break

    if len(selected) < sample_size:
        for row in candidates:
            row_key = (row.query_id, row.precedent_id, row.candidate_rank)
            if row_key in selected_keys:
                continue
            selected.append(row)
            selected_keys.add(row_key)
            if len(selected) >= sample_size:
                break

    return sorted(selected, key=lambda item: (item.query_id, int(item.candidate_rank or "999")))


def model_specific_guidance(model: str) -> str:
    """모델별 평가 성향을 보정하기 위한 일반 지침을 반환한다."""
    normalized = model.lower()
    if "gemma" in normalized:
        return """
추가 지침:
- 세부 당사자, 절차, 금액, 상속 여부가 조금 달라도 질문의 법적 쟁점에 답변 근거로 쓸 수 있으면 2점 이상을 줄 수 있습니다.
- 단순히 사실관계가 완전히 같지 않다는 이유만으로 1점으로 낮추지 마세요.
- 질문이 묻는 결론과 판례가 다루는 쟁점이 연결되는지를 먼저 보세요.
""".strip()
    if "llama" in normalized:
        return """
추가 지침:
- 같은 분야나 비슷한 단어가 있다는 이유만으로 2점 이상을 주지 마세요.
- 질문이 묻는 법적 결론에 직접 도움이 되지 않으면 1점 이하로 평가하세요.
- reason에는 요약을 반복하지 말고 질문과 판례의 일치점 또는 차이점을 직접 비교해 쓰세요.
""".strip()
    return ""


def build_prompt(candidate: CandidateRow, model: str, prompt_mode: str) -> str:
    """로컬 LLM이 후보 판례 관련도를 JSON으로만 평가하도록 프롬프트를 만든다."""
    extra_guidance = ""
    if prompt_mode == "model_specific":
        extra_guidance = "\n\n" + model_specific_guidance(model)
    return f"""
AlphaLawVA 판례 검색 후보를 0~3점으로 평가하세요.
3=질문 상황·사실관계·쟁점이 직접 맞음.
2=사실은 조금 다르지만 같은 쟁점이라 정답 후보 가능.
1=같은 분야이나 직접 답변 근거로는 약함.
0=무관하거나 오해 위험이 큼.
키워드 겹침만으로 높게 주지 말고, 생성요약 내용만 근거로 판단하세요.
상가·영업장·형사·독립 토지거래 중심이면 낮게 보세요.
reason은 질문과 판례를 비교한 판단 근거로 쓰고, 생성요약 문장을 그대로 반복하지 마세요.
JSON만 출력하세요: {{"score": 0, "reason": "한국어 한 문장"}}
{extra_guidance}

[질문]
{candidate.query}

[유형]
{candidate.dispute_type}

[판례]
사건명={candidate.case_name}
사건번호={candidate.case_no}
법원={candidate.court}
요약={candidate.summary}
""".strip()


def result_csv_path(output_dir: Path) -> Path:
    """결과 CSV 경로를 한 곳에서 정한다."""
    return output_dir / "local_llm_candidate_scores.csv"


def extract_json_object(text: str) -> dict[str, Any]:
    """모델 응답에서 JSON 객체를 찾아 파싱한다."""
    stripped = text.strip()
    try:
        return json.loads(stripped)
    except json.JSONDecodeError:
        pass

    start = stripped.find("{")
    end = stripped.rfind("}")
    if start == -1 or end == -1 or end <= start:
        raise ValueError("JSON 객체를 찾지 못했습니다.")
    return json.loads(stripped[start : end + 1])


def call_ollama(
    model: str,
    prompt: str,
    ollama_url: str,
    timeout: int,
    num_ctx: int,
    num_predict: int,
    temperature: float,
) -> tuple[int | None, str, str]:
    """Ollama generate API를 호출하고 score, reason, raw_response를 반환한다."""
    endpoint = ollama_url.rstrip("/") + "/api/generate"
    payload = {
        "model": model,
        "prompt": prompt,
        "stream": False,
        "format": "json",
        "options": {
            "temperature": temperature,
            "num_ctx": num_ctx,
            "num_predict": num_predict,
        },
    }
    request = urllib.request.Request(
        endpoint,
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        body = json.loads(response.read().decode("utf-8"))

    raw_response = compact_text(body.get("response", ""))
    parsed = extract_json_object(raw_response)
    score = int(parsed.get("score"))
    if score not in {0, 1, 2, 3}:
        raise ValueError(f"score가 0~3 범위를 벗어났습니다: {score}")
    reason = compact_text(parsed.get("reason", ""))
    if not reason:
        raise ValueError("reason이 비어 있습니다.")
    return score, reason, raw_response


def result_key(row: dict[str, str]) -> tuple[str, str, str, str]:
    """재실행 시 이미 저장된 모델 판정인지 식별하는 키를 만든다."""
    return (
        row["query_id"],
        row["precedent_id"],
        row["candidate_rank"],
        row["model"],
    )


def load_existing_results(path: Path) -> dict[tuple[str, str, str, str], dict[str, str]]:
    """성공 저장된 CSV 행을 읽어서 재실행 때 건너뛸 수 있게 한다."""
    if not path.exists():
        return {}
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        rows = list(csv.DictReader(file))
    return {result_key(row): row for row in rows if not row.get("error")}


def append_result(path: Path, fieldnames: list[str], row: dict[str, Any]) -> None:
    """모델 판정 결과를 한 건씩 CSV에 즉시 저장한다."""
    exists = path.exists()
    with path.open("a", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        if not exists:
            writer.writeheader()
        writer.writerow(row)


def score_candidates(
    candidates: list[CandidateRow],
    models: list[str],
    args: argparse.Namespace,
    output_dir: Path,
) -> Path:
    """선택된 후보들을 모델별로 평가하고 CSV에 저장한다."""
    result_path = result_csv_path(output_dir)
    fieldnames = [
        "query_id",
        "dispute_type",
        "query",
        "query_specificity",
        "candidate_rank",
        "precedent_id",
        "case_no",
        "case_name",
        "court",
        "decision_date",
        "generated_summary",
        "matched_models",
        "best_rank",
        "matched_chunk_count",
        "codex_score",
        "codex_reason",
        "model",
        "model_score",
        "model_reason",
        "raw_response",
        "error",
        "elapsed_seconds",
        "created_at",
    ]
    existing = load_existing_results(result_path)
    total = len(candidates) * len(models)
    done = 0
    started_at = time.time()

    for candidate in candidates:
        for model in models:
            key = (candidate.query_id, candidate.precedent_id, candidate.candidate_rank, model)
            if key in existing:
                done += 1
                print(f"스킵 {done}/{total}: {candidate.query_id} {candidate.precedent_id} {model}")
                continue

            prompt = build_prompt(candidate, model, args.prompt_mode)
            started = time.time()
            model_score: int | None = None
            model_reason = ""
            raw_response = ""
            error = ""
            try:
                model_score, model_reason, raw_response = call_ollama(
                    model=model,
                    prompt=prompt,
                    ollama_url=args.ollama_url,
                    timeout=args.timeout,
                    num_ctx=args.num_ctx,
                    num_predict=args.num_predict,
                    temperature=args.temperature,
                )
            except (urllib.error.URLError, TimeoutError, ValueError, json.JSONDecodeError) as exc:
                error = str(exc)

            elapsed = round(time.time() - started, 2)
            done += 1
            append_result(
                result_path,
                fieldnames,
                {
                    "query_id": candidate.query_id,
                    "dispute_type": candidate.dispute_type,
                    "query": candidate.query,
                    "query_specificity": candidate.query_specificity,
                    "candidate_rank": candidate.candidate_rank,
                    "precedent_id": candidate.precedent_id,
                    "case_no": candidate.case_no,
                    "case_name": candidate.case_name,
                    "court": candidate.court,
                    "decision_date": candidate.decision_date,
                    "generated_summary": candidate.summary,
                    "matched_models": candidate.matched_models,
                    "best_rank": candidate.best_rank,
                    "matched_chunk_count": candidate.matched_chunk_count,
                    "codex_score": candidate.codex_score,
                    "codex_reason": candidate.codex_reason,
                    "model": model,
                    "model_score": "" if model_score is None else model_score,
                    "model_reason": model_reason,
                    "raw_response": raw_response,
                    "error": error,
                    "elapsed_seconds": elapsed,
                    "created_at": now_utc_iso(),
                },
            )
            status = "실패" if error else f"{model_score}점"
            average = (time.time() - started_at) / done
            remaining = max(total - done, 0) * average
            print(
                f"[{done}/{total}] {model} {candidate.query_id} "
                f"후보{candidate.candidate_rank} {status} "
                f"{elapsed:.1f}초 예상남음 {remaining:.0f}초",
                flush=True,
            )
    return result_path


def read_result_rows(result_path: Path) -> list[dict[str, str]]:
    """CSV 결과를 읽어 엑셀 시트 작성용 row 목록으로 반환한다."""
    with result_path.open("r", encoding="utf-8-sig", newline="") as file:
        return list(csv.DictReader(file))


def build_wide_rows(result_rows: list[dict[str, str]], models: list[str]) -> list[dict[str, str]]:
    """모델별 세로 결과를 후보 판례 1행당 모델 컬럼이 붙는 넓은 형태로 바꾼다."""
    grouped: dict[tuple[str, str, str], dict[str, str]] = {}
    for row in result_rows:
        key = (row["query_id"], row["precedent_id"], row["candidate_rank"])
        base = grouped.setdefault(
            key,
            {
                "query_id": row["query_id"],
                "dispute_type": row["dispute_type"],
                "query": row["query"],
                "query_specificity": row["query_specificity"],
                "candidate_rank": row["candidate_rank"],
                "precedent_id": row["precedent_id"],
                "case_no": row["case_no"],
                "case_name": row["case_name"],
                "court": row["court"],
                "decision_date": row["decision_date"],
                "generated_summary": row["generated_summary"],
                "matched_models": row["matched_models"],
                "best_rank": row["best_rank"],
                "matched_chunk_count": row["matched_chunk_count"],
                "codex_score": row["codex_score"],
                "codex_reason": row["codex_reason"],
            },
        )
        model_key = row["model"]
        base[f"{model_key}_score"] = row["model_score"]
        base[f"{model_key}_reason"] = row["model_reason"]
        base[f"{model_key}_error"] = row["error"]

    for row in grouped.values():
        scores = [row.get("codex_score", "")]
        scores.extend(row.get(f"{model}_score", "") for model in models)
        parsed_scores = [score for score in scores if str(score).strip() != ""]
        row["all_scores"] = " / ".join(str(score) for score in scores)
        row["score_agreement"] = "일치" if len(set(parsed_scores)) == 1 else "불일치"
        strong_scores = [int(score) for score in parsed_scores if str(score).isdigit()]
        row["review_priority"] = (
            "높음"
            if len(set(str(score) for score in parsed_scores)) > 1
            or any(score >= 2 for score in strong_scores)
            else "낮음"
        )
    return sorted(grouped.values(), key=lambda row: (row["query_id"], int(row["candidate_rank"])))


def write_workbook_sheet(
    workbook_path: Path,
    sheet_name: str,
    wide_rows: list[dict[str, str]],
    models: list[str],
) -> None:
    """모델 비교 결과를 워크북의 새 시트로 저장한다."""
    workbook = load_workbook(workbook_path, keep_vba=True)
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    sheet = workbook.create_sheet(sheet_name)

    headers = [
        "질문ID",
        "분쟁유형",
        "질문",
        "생성요약",
        "질문구체성",
        "후보순위",
        "판례일련번호",
        "사건번호",
        "사건명",
        "법원",
        "선고일자",
        "검색모델",
        "최고순위",
        "검색청크수",
        "기존AI점수",
        "기존AI근거",
    ]
    for model in models:
        headers.extend([f"{model}_점수", f"{model}_근거", f"{model}_오류"])
    headers.extend(["점수모음", "점수일치여부", "검토우선순위"])

    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrap_top = Alignment(vertical="top", wrap_text=True)

    for col_index, value in enumerate(headers, 1):
        cell = sheet.cell(1, col_index, value)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for row_index, row in enumerate(wide_rows, 2):
        values = [
            row["query_id"],
            row["dispute_type"],
            row["query"],
            row["generated_summary"],
            row["query_specificity"],
            row["candidate_rank"],
            row["precedent_id"],
            row["case_no"],
            row["case_name"],
            row["court"],
            row["decision_date"],
            row["matched_models"],
            row["best_rank"],
            row["matched_chunk_count"],
            row["codex_score"],
            row["codex_reason"],
        ]
        for model in models:
            values.extend(
                [
                    row.get(f"{model}_score", ""),
                    row.get(f"{model}_reason", ""),
                    row.get(f"{model}_error", ""),
                ]
            )
        values.extend([row["all_scores"], row["score_agreement"], row["review_priority"]])

        for col_index, value in enumerate(values, 1):
            cell = sheet.cell(row_index, col_index, value)
            cell.alignment = wrap_top if col_index in {2, 3, 4, 9, 16, 18, 21} else center
            cell.border = border

    widths = {
        1: 20,
        2: 24,
        3: 52,
        4: 62,
        5: 16,
        6: 10,
        7: 14,
        8: 20,
        9: 24,
        10: 18,
        11: 14,
        12: 18,
        13: 10,
        14: 12,
        15: 12,
        16: 44,
        17: 12,
        18: 44,
        19: 24,
        20: 12,
        21: 44,
        22: 24,
        23: 16,
        24: 14,
        25: 14,
    }
    for col_index, width in widths.items():
        sheet.column_dimensions[get_column_letter(col_index)].width = width
    for row_index in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_index].height = 82

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(workbook_path)


def write_manifest(
    output_dir: Path,
    args: argparse.Namespace,
    models: list[str],
    selected_candidates: list[CandidateRow],
    result_rows: list[dict[str, str]],
) -> None:
    """실험 재현을 위한 설정과 간단한 통계를 manifest로 저장한다."""
    score_counts: dict[str, dict[str, int]] = {}
    for model in models:
        counter = Counter(
            row["model_score"]
            for row in result_rows
            if row["model"] == model and row["model_score"] != ""
        )
        score_counts[model] = dict(sorted(counter.items()))

    manifest = {
        "schema_version": "local_llm_candidate_score_compare.v1",
        "created_at": now_utc_iso(),
        "workbook": str(args.workbook),
        "candidate_sheet": args.candidate_sheet,
        "result_sheet": args.result_sheet,
        "models": models,
        "prompt_mode": args.prompt_mode,
        "sample_size": len(selected_candidates),
        "total_model_calls": len(selected_candidates) * len(models),
        "score_counts": score_counts,
        "prompt_rule": "question + dispute_type + case metadata + generated_summary only",
    }
    (output_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def main() -> None:
    """샘플 후보 판례를 로컬 LLM으로 평가하고 결과 파일과 시트를 만든다."""
    args = parse_args()
    models = [model.strip() for model in args.models.split(",") if model.strip()]
    if not models:
        raise ValueError("비교할 모델이 없습니다.")

    output_dir = make_output_dir(args.output_dir)
    candidates = read_candidates(args.workbook, args.candidate_sheet)
    selected_candidates = select_sample(candidates, args.sample_size)

    print(f"로컬 LLM 후보 점수 비교 시작: 후보 {len(selected_candidates)}건, 모델 {len(models)}개")
    print(f"모델: {', '.join(models)}")
    print(f"결과 폴더: {output_dir}")

    try:
        result_path = score_candidates(selected_candidates, models, args, output_dir)
    except KeyboardInterrupt:
        result_path = result_csv_path(output_dir)
        print("\n중단됨: 이미 저장된 CSV가 있으면 그 결과만으로 시트를 갱신합니다.")
        if not result_path.exists():
            raise
    result_rows = read_result_rows(result_path)
    wide_rows = build_wide_rows(result_rows, models)

    if not args.skip_workbook:
        write_workbook_sheet(args.workbook, args.result_sheet, wide_rows, models)

    write_manifest(output_dir, args, models, selected_candidates, result_rows)
    print("완료")
    print(f"- CSV: {result_path}")
    print(f"- manifest: {output_dir / 'manifest.json'}")
    if not args.skip_workbook:
        print(f"- workbook sheet: {args.result_sheet}")


if __name__ == "__main__":
    main()
