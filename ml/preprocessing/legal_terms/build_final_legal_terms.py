# build_final_legal_terms.py
"""
Description: 사용자 난이도 기준으로 선별된 법률 분야 용어와 수동 선별 용어를
원 정의·출처 메타데이터를 보존한 최종 용어 후보 파일로 통합한다.
Author: choeminju
Date: 2026-10-05
Before:
    - 법률 분야 후보선택 결과와 수동 검토 엑셀의 바로 적재 후보가 준비된 상태.
After:
    - 최종 981개 용어 JSONL·CSV와 쉬운 정의 재실험용 30개 표본이 생성.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable


PROJECT_ROOT = Path(__file__).resolve().parents[3]
LOCAL_DATA_ROOT = PROJECT_ROOT / "local_data" / "legal_terms"
DEFAULT_LEGAL_PATH = (
    LOCAL_DATA_ROOT
    / "reviews"
    / "user_facing_legal_terms_v01"
    / "selected_terms.jsonl"
)
DEFAULT_MANUAL_REVIEW_PATH = (
    LOCAL_DATA_ROOT / "reviews" / "legal_term_candidate_ai_review_v01.xlsx"
)
DEFAULT_FULL_SOURCE_PATH = LOCAL_DATA_ROOT / "integrated_v01" / "usable_legal_terms.jsonl"
DEFAULT_COMPONENT_SOURCE_PATH = (
    LOCAL_DATA_ROOT
    / "urimalsaem_component_definitions_v01"
    / "urimalsaem_term_definitions.jsonl"
)
DEFAULT_OUTPUT_DIR = LOCAL_DATA_ROOT / "final_v01"

MULTI_SAMPLE_TERMS = [
    "법제",
    "경합",
    "공소",
    "각하",
    "보호처분",
    "중재계약",
    "직권조사",
    "담보책임",
    "처분행위",
    "방조",
    "조건부권리",
    "촉탁",
    "보전",
    "본등기",
    "비채변제",
    "통행권",
    "소송계속",
    "수임",
    "이심",
    "추완",
]
SINGLE_SAMPLE_TERMS = [
    "가등기가처분",
    "가액반환",
    "거소지정권",
    "납입담보책임",
    "대항력",
    "명의수탁자",
    "명의신탁약정",
    "무자력",
    "석명",
    "선관의무",
]


def parse_args() -> argparse.Namespace:
    """통합 입력과 출력 경로를 정의한다."""
    parser = argparse.ArgumentParser(description="최종 법률 용어 후보 통합")
    parser.add_argument("--legal-path", type=Path, default=DEFAULT_LEGAL_PATH)
    parser.add_argument(
        "--manual-review-path", type=Path, default=DEFAULT_MANUAL_REVIEW_PATH
    )
    parser.add_argument("--full-source-path", type=Path, default=DEFAULT_FULL_SOURCE_PATH)
    parser.add_argument(
        "--component-source-path", type=Path, default=DEFAULT_COMPONENT_SOURCE_PATH
    )
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    return parser.parse_args()


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    """JSONL 파일을 행 목록으로 읽는다."""
    with path.open("r", encoding="utf-8-sig") as file:
        return [json.loads(line) for line in file if line.strip()]


def write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    """행 목록을 UTF-8 JSONL로 저장한다."""
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_suffix(f"{path.suffix}.tmp")
    with temporary_path.open("w", encoding="utf-8") as file:
        for row in rows:
            file.write(json.dumps(row, ensure_ascii=False) + "\n")
    temporary_path.replace(path)


def stable_definition_id(definition: dict[str, Any]) -> str:
    """우리말샘 식별자가 없을 때도 안정적인 정의 ID를 만든다."""
    target_code = str(definition.get("target_code", "")).strip()
    if target_code:
        return target_code
    payload = "|".join(
        [
            str(definition.get("headword", "")),
            str(definition.get("sense_no", "")),
            str(definition.get("definition", "")),
        ]
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:16]


def normalize_definition(definition: dict[str, Any]) -> dict[str, Any]:
    """서로 다른 입력의 우리말샘 정의를 공통 구조로 정규화한다."""
    return {
        "definition_id": stable_definition_id(definition),
        "provider": str(definition.get("provider", "urimalsaem")),
        "lookup_term": str(definition.get("lookup_term", "")),
        "headword": str(definition.get("headword", "")),
        "sense_no": str(definition.get("sense_no", "")),
        "category": str(definition.get("category", "")),
        "source_definition": str(definition.get("definition", "")).strip(),
        "part_of_speech": str(definition.get("part_of_speech", "")),
        "origin": str(definition.get("origin", "")),
        "source_link": str(definition.get("source_link", "")),
    }


def source_stats(row: dict[str, Any]) -> dict[str, int]:
    """판례·법령 출처 통계를 공통 구조로 반환한다."""
    return {
        "precedent_case_count": int(row.get("precedent_case_count") or 0),
        "precedent_occurrence_count": int(row.get("precedent_occurrence_count") or 0),
        "statute_document_count": int(row.get("statute_document_count") or 0),
        "statute_article_count": int(row.get("statute_article_count") or 0),
        "statute_occurrence_count": int(row.get("statute_occurrence_count") or 0),
    }


def normalize_legal_row(row: dict[str, Any]) -> dict[str, Any]:
    """법률 분야 후보선택 행을 최종 스키마로 변환한다."""
    definitions = [normalize_definition(item) for item in row.get("legal_definitions", [])]
    return {
        "schema_version": "legal_term_glossary_candidate.v0.1",
        "match_key": str(row["match_key"]),
        "term": str(row["term"]),
        "source_domains": list(row.get("source_domains") or []),
        "selection_source": "legal_category_user_difficulty_review",
        "selection_reason": str(row.get("selection_reason", "")),
        "generation_status": str(row.get("generation_status", "ready")),
        "definition_warning": str(row.get("definition_warning", "")),
        "definitions": definitions,
        "definition_count": len(definitions),
        "source_stats": source_stats(row),
    }


def load_manual_review_rows(path: Path) -> list[dict[str, Any]]:
    """사용자 후보선택 중 바로 적재 권고를 받은 142개 행을 읽는다."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("수동 검토 XLSX를 읽으려면 openpyxl이 필요합니다.") from exc

    workbook = load_workbook(path, data_only=True, read_only=False)
    worksheet = workbook["후보전수검토"]
    headers = [cell.value for cell in worksheet[1]]
    rows = []
    for cells in worksheet.iter_rows(min_row=2):
        row = dict(zip(headers, [cell.value for cell in cells]))
        if row.get("사용자 검토결정") == "후보선택" and row.get("AI판정") == "유지":
            rows.append(row)
    return rows


def build_dictionary_index(
    full_rows: list[dict[str, Any]], component_rows: list[dict[str, Any]]
) -> dict[str, list[dict[str, Any]]]:
    """전체 용어와 구성 용어의 우리말샘 정의를 매치키별로 색인한다."""
    index: dict[str, list[dict[str, Any]]] = {}
    for row in full_rows:
        index[str(row["match_key"])] = list(
            (row.get("definitions") or {}).get("urimalsaem", [])
        )
    for row in component_rows:
        index[str(row["match_key"])] = list(row.get("definitions") or [])
    return index


def normalize_manual_row(
    row: dict[str, Any], dictionary_index: dict[str, list[dict[str, Any]]]
) -> dict[str, Any]:
    """수동 검토에서 선택한 뜻을 원본 우리말샘 정의와 다시 연결한다."""
    match_key = str(row["매치키"])
    selected_text = str(row["AI추천 뜻풀이"]).strip()
    matches = [
        definition
        for definition in dictionary_index.get(match_key, [])
        if str(definition.get("definition", "")).strip() == selected_text
    ]
    if len(matches) != 1:
        raise ValueError(
            f"{match_key}의 선택 정의를 원본과 1:1로 연결할 수 없습니다: {len(matches)}개"
        )
    definition = normalize_definition(matches[0])
    return {
        "schema_version": "legal_term_glossary_candidate.v0.1",
        "match_key": match_key,
        "term": str(row["용어"]),
        "source_domains": [str(row.get("출처영역") or "")],
        "selection_source": "manual_nonlegal_definition_review",
        "selection_reason": "사용자 후보선택 후 바로 적재 권고",
        "generation_status": "ready",
        "definition_warning": "",
        "definitions": [definition],
        "definition_count": 1,
        "source_stats": {
            "precedent_case_count": int(row.get("판례건수") or 0),
            "precedent_occurrence_count": int(row.get("판례출현수") or 0),
            "statute_document_count": int(row.get("법령수") or 0),
            "statute_article_count": int(row.get("조문수") or 0),
            "statute_occurrence_count": int(row.get("법령출현수") or 0),
        },
    }


def merge_rows(
    legal_rows: list[dict[str, Any]], manual_rows: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    """두 후보군을 합치고 매치키 중복을 차단한다."""
    merged = legal_rows + manual_rows
    keys = [row["match_key"] for row in merged]
    duplicates = sorted(key for key, count in Counter(keys).items() if count > 1)
    if duplicates:
        raise ValueError(f"통합 후보에 중복 매치키가 있습니다: {duplicates[:10]}")
    return sorted(merged, key=lambda row: row["match_key"])


def select_sample(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """어려운 다중 정의 20개와 단일 정의 10개를 고정 표본으로 선택한다."""
    by_term = {row["term"]: row for row in rows}
    terms = MULTI_SAMPLE_TERMS + SINGLE_SAMPLE_TERMS
    missing = [term for term in terms if term not in by_term]
    if missing:
        raise ValueError(f"표본 용어가 최종 후보에 없습니다: {missing}")
    sample = [by_term[term] for term in terms]
    multi_count = sum(row["definition_count"] >= 2 for row in sample)
    single_count = sum(row["definition_count"] == 1 for row in sample)
    if multi_count != 20 or single_count != 10:
        raise ValueError(
            f"표본 정의 구성 오류: 다중 {multi_count}개, 단일 {single_count}개"
        )
    return sample


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    """최종 후보를 한 용어 한 행의 검토용 CSV로 저장한다."""
    headers = [
        "term",
        "match_key",
        "selection_source",
        "generation_status",
        "definition_warning",
        "definition_count",
        "definition_categories",
        "source_definitions",
        "precedent_case_count",
        "statute_document_count",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    "term": row["term"],
                    "match_key": row["match_key"],
                    "selection_source": row["selection_source"],
                    "generation_status": row["generation_status"],
                    "definition_warning": row["definition_warning"],
                    "definition_count": row["definition_count"],
                    "definition_categories": " | ".join(
                        item["category"] for item in row["definitions"]
                    ),
                    "source_definitions": "\n".join(
                        f"{index}. {item['source_definition']}"
                        for index, item in enumerate(row["definitions"], start=1)
                    ),
                    "precedent_case_count": row["source_stats"][
                        "precedent_case_count"
                    ],
                    "statute_document_count": row["source_stats"][
                        "statute_document_count"
                    ],
                }
            )


def main() -> None:
    """두 후보군을 통합하고 최종 파일·표본·manifest를 저장한다."""
    args = parse_args()
    legal_rows = [normalize_legal_row(row) for row in read_jsonl(args.legal_path)]
    dictionary_index = build_dictionary_index(
        read_jsonl(args.full_source_path),
        read_jsonl(args.component_source_path),
    )
    manual_rows = [
        normalize_manual_row(row, dictionary_index)
        for row in load_manual_review_rows(args.manual_review_path)
    ]
    merged = merge_rows(legal_rows, manual_rows)
    if len(merged) != 981:
        raise ValueError(f"최종 후보 수가 981개가 아닙니다: {len(merged)}개")
    if any(not row["definitions"] for row in merged):
        raise ValueError("원 정의가 없는 최종 후보가 있습니다.")

    sample = select_sample(merged)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_jsonl(args.output_dir / "legal_terms_981.jsonl", merged)
    write_csv(args.output_dir / "legal_terms_981.csv", merged)
    write_jsonl(args.output_dir / "easy_definition_sample30.jsonl", sample)

    manifest = {
        "schema_version": "legal_term_glossary_build.v0.1",
        "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "total_term_count": len(merged),
        "total_definition_count": sum(row["definition_count"] for row in merged),
        "selection_source_counts": dict(
            Counter(row["selection_source"] for row in merged)
        ),
        "generation_status_counts": dict(
            Counter(row["generation_status"] for row in merged)
        ),
        "sample_term_count": len(sample),
        "sample_multi_definition_count": sum(
            row["definition_count"] >= 2 for row in sample
        ),
        "sample_single_definition_count": sum(
            row["definition_count"] == 1 for row in sample
        ),
        "sample_terms": [row["term"] for row in sample],
    }
    (args.output_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"최종 용어: {len(merged)}개")
    print(f"최종 원 정의: {manifest['total_definition_count']}개")
    print("표본: 다중 정의 20개 + 단일 정의 10개")
    print(f"결과 폴더: {args.output_dir}")


if __name__ == "__main__":
    main()
