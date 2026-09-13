# pool_precedent_candidates.py
"""
Description: 판례 검색 평가셋 질문을 3개 임베딩 벡터DB에 검색해
판례 단위 후보 풀과 검색 청크 목록을 생성한다.
Author: choeminju
Date: 2026-09-06
Before:
    - A_reason_summary_v1 청크로 만든 Chroma 벡터DB 3개가 준비된 상태.
    - 평가셋 질문 후보 파일에 query_id와 query_candidate가 정리된 상태.
After:
    - 질문별 후보 판례 CSV와 후보별 검색 청크 CSV가 생성된다.
    - 같은 판례에서 검색된 여러 청크는 precedent_id 기준으로 묶인다.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import time
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


# HuggingFace tokenizer가 Chroma 내부 처리와 함께 쓰일 때 종료 단계에서
# 병렬 처리 경고가 반복될 수 있어 평가 스크립트에서는 기본적으로 끈다.
os.environ.setdefault("TOKENIZERS_PARALLELISM", "false")

PROJECT_ROOT = Path(__file__).resolve().parents[3]
LOCAL_DATA_ROOT = PROJECT_ROOT / "local_data"
DEFAULT_VECTOR_DB_ROOT = LOCAL_DATA_ROOT / "precedents" / "vector_dbs"
DEFAULT_FINAL_CASES_DIR = (
    LOCAL_DATA_ROOT / "precedents" / "processed" / "final_cases"
)
DEFAULT_OUTPUT_ROOT = LOCAL_DATA_ROOT / "precedents" / "evaluation" / "pooling_runs"
DEFAULT_TOP_K = 20
DEFAULT_PREVIEW_CHARS = 500

MODEL_CONFIGS = {
    "kure": {
        "provider": "sentence_transformers",
        "model": "nlpai-lab/KURE-v1",
        "db_dir": "A_reason_summary_v1__kure",
    },
    "bge-m3": {
        "provider": "sentence_transformers",
        "model": "BAAI/bge-m3",
        "db_dir": "A_reason_summary_v1__bge-m3",
    },
    "openai": {
        "provider": "openai",
        "model": "text-embedding-3-large",
        "db_dir": "A_reason_summary_v1__openai",
    },
}


@dataclass(slots=True)
class QueryItem:
    """평가셋 질문 한 건을 담는 내부 구조."""

    query_id: str
    query_candidate: str
    dispute_type: str = ""
    query_difficulty: str = ""
    scenario_tags: str = ""


@dataclass(slots=True)
class MatchedChunk:
    """벡터DB에서 검색된 청크 한 건을 담는 내부 구조."""

    model_key: str
    rank: int
    chunk_id: str
    chunk_type: str
    distance: float | None
    text: str


@dataclass(slots=True)
class CandidateCase:
    """청크 검색 결과를 판례 단위로 묶은 후보 구조."""

    precedent_id: str
    chunks: list[MatchedChunk] = field(default_factory=list)
    best_rank: int = 10**9
    best_distance: float | None = None

    def add_chunk(self, chunk: MatchedChunk) -> None:
        """후보 판례에 검색 청크를 추가하고 최상위 순위를 갱신한다."""
        self.chunks.append(chunk)
        self.best_rank = min(self.best_rank, chunk.rank)
        if chunk.distance is None:
            return
        if self.best_distance is None or chunk.distance < self.best_distance:
            self.best_distance = chunk.distance

    @property
    def matched_models(self) -> list[str]:
        """이 판례를 검색한 임베딩 모델 목록을 정렬해 반환한다."""
        return sorted({chunk.model_key for chunk in self.chunks})


def parse_args() -> argparse.Namespace:
    """명령행 옵션을 정의한다."""
    parser = argparse.ArgumentParser(
        description="Pool candidate precedents from multiple vector DBs.",
    )
    input_group = parser.add_mutually_exclusive_group(required=True)
    input_group.add_argument(
        "--queries-csv",
        type=Path,
        help="평가 질문 CSV 경로. query_id, query_candidate 컬럼이 필요하다.",
    )
    input_group.add_argument(
        "--queries-xlsx",
        type=Path,
        help="평가 질문 XLSX/XLSM 경로. Google Sheets에서 받은 파일도 가능하다.",
    )
    parser.add_argument(
        "--sheet-name",
        default="질문후보_v2",
        help="XLSX/XLSM에서 읽을 시트 이름.",
    )
    parser.add_argument(
        "--selected-values",
        default="선택",
        help="평가 컬럼에서 포함할 값. 여러 개면 쉼표로 구분한다.",
    )
    parser.add_argument(
        "--query-limit",
        type=int,
        default=None,
        help="테스트용 질문 개수 제한.",
    )
    parser.add_argument(
        "--one-per-dispute-type",
        action="store_true",
        help="분쟁유형별로 최대 1개씩만 뽑아 query-limit까지 사용한다.",
    )
    parser.add_argument(
        "--top-k",
        type=int,
        default=DEFAULT_TOP_K,
        help="모델별 검색할 청크 개수.",
    )
    parser.add_argument(
        "--embeddings",
        default="kure,bge-m3,openai",
        help="검색할 임베딩 설정. 예: kure,bge-m3,openai",
    )
    parser.add_argument(
        "--vector-db-root",
        type=Path,
        default=DEFAULT_VECTOR_DB_ROOT,
        help="A_reason_summary_v1__* 벡터DB 폴더들이 들어 있는 루트.",
    )
    parser.add_argument(
        "--final-cases-dir",
        type=Path,
        default=DEFAULT_FINAL_CASES_DIR,
        help="생성요약이 포함된 최종 판례 JSON 폴더.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=None,
        help="결과 저장 폴더. 기본값은 local_data 아래 타임스탬프 폴더.",
    )
    parser.add_argument(
        "--preview-chars",
        type=int,
        default=DEFAULT_PREVIEW_CHARS,
        help="검색 청크 미리보기 글자 수.",
    )
    return parser.parse_args()


def load_dotenv(path: Path = PROJECT_ROOT / ".env") -> None:
    """로컬 .env 파일의 KEY=VALUE 값을 환경변수로 불러온다."""
    if not path.exists():
        return

    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        value = value.strip().strip('"').strip("'")
        os.environ.setdefault(key.strip(), value)


def now_utc_iso() -> str:
    """현재 UTC 시각을 ISO 문자열로 반환한다."""
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def compact_text(value: Any) -> str:
    """CSV에 넣기 좋게 줄바꿈과 연속 공백을 정리한다."""
    if value is None:
        return ""
    return " ".join(str(value).split())


def parse_embedding_keys(raw_value: str) -> list[str]:
    """쉼표로 받은 임베딩 설정 이름을 검증해 목록으로 반환한다."""
    keys = [value.strip() for value in raw_value.split(",") if value.strip()]
    unknown = [key for key in keys if key not in MODEL_CONFIGS]
    if unknown:
        raise ValueError(f"지원하지 않는 임베딩 설정입니다: {', '.join(unknown)}")
    return keys


def load_queries_from_csv(path: Path, selected_values: set[str]) -> list[QueryItem]:
    """CSV 파일에서 선택된 평가 질문을 읽는다."""
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        rows = list(csv.DictReader(file))
    return rows_to_queries(rows, selected_values)


def load_queries_from_xlsx(
    path: Path,
    sheet_name: str,
    selected_values: set[str],
) -> list[QueryItem]:
    """XLSX/XLSM 파일에서 선택된 평가 질문을 읽는다."""
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "XLSX/XLSM 입력을 읽으려면 openpyxl이 필요합니다. "
            "CSV로 내보내서 --queries-csv를 쓰거나 openpyxl을 설치해주세요."
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"시트를 찾지 못했습니다: {sheet_name}")

    sheet = workbook[sheet_name]
    header_row_index: int | None = None
    headers: list[str] = []

    for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = [compact_text(value) for value in row]
        if "query_id" in values and "query_candidate" in values:
            header_row_index = row_index
            headers = values
            break

    if header_row_index is None:
        raise ValueError("query_id와 query_candidate가 있는 헤더 행을 찾지 못했습니다.")

    rows: list[dict[str, Any]] = []
    for row in sheet.iter_rows(min_row=header_row_index + 1, values_only=True):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        rows.append(dict(zip(headers, row, strict=False)))

    return rows_to_queries(rows, selected_values)


def rows_to_queries(rows: list[dict[str, Any]], selected_values: set[str]) -> list[QueryItem]:
    """표 형태 row 목록을 QueryItem 목록으로 변환한다."""
    queries: list[QueryItem] = []
    for row in rows:
        review_value = compact_text(row.get("평가"))
        query_id = compact_text(row.get("query_id"))
        query_candidate = compact_text(row.get("query_candidate"))
        if selected_values and review_value not in selected_values:
            continue
        if not query_id or not query_candidate:
            continue
        queries.append(
            QueryItem(
                query_id=query_id,
                query_candidate=query_candidate,
                dispute_type=compact_text(row.get("dispute_type")),
                query_difficulty=compact_text(row.get("query_difficulty")),
                scenario_tags=compact_text(row.get("scenario_tags")),
            )
        )
    return queries


def select_queries(
    queries: list[QueryItem],
    query_limit: int | None,
    one_per_dispute_type: bool,
) -> list[QueryItem]:
    """테스트 옵션에 맞게 질문 목록을 줄인다."""
    if query_limit is None:
        return queries

    if not one_per_dispute_type:
        return queries[:query_limit]

    selected: list[QueryItem] = []
    seen_dispute_types: set[str] = set()
    for query in queries:
        dispute_type = query.dispute_type or query.query_id
        if dispute_type in seen_dispute_types:
            continue
        selected.append(query)
        seen_dispute_types.add(dispute_type)
        if len(selected) >= query_limit:
            break
    return selected


def read_json(path: Path) -> dict[str, Any]:
    """JSON 파일 하나를 dict로 읽는다."""
    return json.loads(path.read_text(encoding="utf-8"))


def load_vector_db_manifest(vector_db_root: Path, embedding_key: str) -> dict[str, Any]:
    """벡터DB 폴더의 manifest에서 컬렉션명과 모델 정보를 읽는다."""
    db_dir = vector_db_root / MODEL_CONFIGS[embedding_key]["db_dir"]
    manifest_path = db_dir / "manifest.json"
    if not manifest_path.exists():
        raise FileNotFoundError(f"벡터DB manifest를 찾지 못했습니다: {manifest_path}")
    manifest = read_json(manifest_path)
    manifest["db_dir"] = str(db_dir)
    return manifest


def load_sentence_transformer(model_name: str) -> Any:
    """SentenceTransformers 임베딩 모델을 로드한다."""
    try:
        from sentence_transformers import SentenceTransformer
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "KURE/BGE-M3 검색에는 sentence-transformers가 필요합니다. "
            "설치 뒤 다시 실행해주세요."
        ) from exc
    return SentenceTransformer(model_name, trust_remote_code=True)


def load_openai_client() -> Any:
    """OpenAI 임베딩 API 클라이언트를 생성한다."""
    try:
        from openai import OpenAI
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "OpenAI 임베딩 검색에는 openai 패키지가 필요합니다. "
            "설치 뒤 다시 실행해주세요."
        ) from exc
    return OpenAI(api_key=os.environ.get("OPENAI_API_KEY"))


def embed_query(provider: str, model: Any, model_name: str, query: str) -> list[float]:
    """사용자 질문 한 개를 해당 벡터DB와 같은 임베딩 모델로 변환한다."""
    if provider == "sentence_transformers":
        vector = model.encode([query], normalize_embeddings=True, show_progress_bar=False)
        return vector.tolist()[0]

    if provider == "openai":
        response = model.embeddings.create(model=model_name, input=[query])
        return response.data[0].embedding

    raise ValueError(f"지원하지 않는 provider입니다: {provider}")


def load_chroma_collection(vector_db_root: Path, embedding_key: str) -> tuple[Any, dict[str, Any]]:
    """Chroma 벡터DB를 열고 manifest에 기록된 컬렉션을 반환한다."""
    try:
        import chromadb
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "벡터DB 검색에는 chromadb가 필요합니다. 설치 뒤 다시 실행해주세요."
        ) from exc

    manifest = load_vector_db_manifest(vector_db_root, embedding_key)
    db_dir = Path(manifest["db_dir"])
    client = chromadb.PersistentClient(path=str(db_dir / "chroma"))
    return client.get_collection(manifest["collection"]), manifest


def search_collection(
    collection: Any,
    query_vector: list[float],
    top_k: int,
) -> dict[str, Any]:
    """임베딩된 질문 벡터로 Chroma 컬렉션에서 Top-K 청크를 검색한다."""
    return collection.query(
        query_embeddings=[query_vector],
        n_results=top_k,
        include=["documents", "metadatas", "distances"],
    )


def load_final_case(final_cases_dir: Path, precedent_id: str) -> dict[str, Any]:
    """후보 판례의 final_cases JSON을 읽어 기본정보와 생성요약을 가져온다."""
    path = final_cases_dir / f"{precedent_id}.json"
    if not path.exists():
        return {}
    return read_json(path)


def build_candidate_pool(
    query: QueryItem,
    embedding_key: str,
    search_result: dict[str, Any],
    pooled: dict[str, CandidateCase],
) -> None:
    """검색 청크 결과를 precedent_id 기준 후보 풀에 합친다."""
    ids = search_result.get("ids", [[]])[0]
    documents = search_result.get("documents", [[]])[0]
    metadatas = search_result.get("metadatas", [[]])[0]
    distances = search_result.get("distances", [[]])[0]

    for index, chunk_id in enumerate(ids):
        metadata = metadatas[index] or {}
        precedent_id = compact_text(
            metadata.get("precedent_id") or metadata.get("source_case_id")
        )
        if not precedent_id:
            continue

        candidate = pooled.setdefault(
            precedent_id,
            CandidateCase(precedent_id=precedent_id),
        )
        candidate.add_chunk(
            MatchedChunk(
                model_key=embedding_key,
                rank=index + 1,
                chunk_id=compact_text(chunk_id),
                chunk_type=compact_text(metadata.get("section") or metadata.get("chunk_type")),
                distance=distances[index] if index < len(distances) else None,
                text=documents[index] if index < len(documents) else "",
            )
        )


def sort_candidates(candidates: dict[str, CandidateCase]) -> list[CandidateCase]:
    """후보 판례를 검수하기 좋은 순서로 정렬한다."""
    return sorted(
        candidates.values(),
        key=lambda item: (
            item.best_rank,
            -len(item.matched_models),
            item.best_distance if item.best_distance is not None else 10**9,
            item.precedent_id,
        ),
    )


def make_output_dir(base_dir: Path | None) -> Path:
    """결과 저장 폴더를 만들고 반환한다."""
    if base_dir is not None:
        output_dir = base_dir
    else:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_dir = DEFAULT_OUTPUT_ROOT / f"A_reason_summary_v1_topk_{timestamp}"
    output_dir.mkdir(parents=True, exist_ok=True)
    return output_dir


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    """dict row 목록을 CSV 파일로 저장한다."""
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def build_output_rows(
    query: QueryItem,
    candidates: list[CandidateCase],
    final_cases_dir: Path,
    preview_chars: int,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """판례 후보 CSV row와 검색 청크 CSV row를 생성한다."""
    candidate_rows: list[dict[str, Any]] = []
    chunk_rows: list[dict[str, Any]] = []

    for candidate_rank, candidate in enumerate(candidates, start=1):
        final_case = load_final_case(final_cases_dir, candidate.precedent_id)
        matched_models = candidate.matched_models

        candidate_rows.append(
            {
                "query_id": query.query_id,
                "dispute_type": query.dispute_type,
                "query_candidate": query.query_candidate,
                "query_difficulty": query.query_difficulty,
                "scenario_tags": query.scenario_tags,
                "candidate_rank": candidate_rank,
                "precedent_id": candidate.precedent_id,
                "case_no": compact_text(final_case.get("사건번호")),
                "case_name": compact_text(final_case.get("사건명")),
                "court": compact_text(final_case.get("법원명")),
                "decision_date": compact_text(final_case.get("선고일자")),
                "generated_summary": compact_text(final_case.get("생성요약")),
                "matched_models": "|".join(matched_models),
                "best_rank": candidate.best_rank,
                "best_distance": candidate.best_distance,
                "matched_chunk_count": len(candidate.chunks),
                "gold_label": "",
                "label_reason": "",
            }
        )

        for chunk in sorted(candidate.chunks, key=lambda item: (item.model_key, item.rank)):
            chunk_rows.append(
                {
                    "query_id": query.query_id,
                    "precedent_id": candidate.precedent_id,
                    "candidate_rank": candidate_rank,
                    "model": chunk.model_key,
                    "rank": chunk.rank,
                    "chunk_id": chunk.chunk_id,
                    "chunk_type": chunk.chunk_type,
                    "distance": chunk.distance,
                    "chunk_text_preview": compact_text(chunk.text)[:preview_chars],
                }
            )

    return candidate_rows, chunk_rows


def main() -> None:
    """평가 질문을 검색해 판례 후보 풀링 결과를 저장한다."""
    args = parse_args()
    load_dotenv()

    selected_values = {
        value.strip() for value in args.selected_values.split(",") if value.strip()
    }
    embedding_keys = parse_embedding_keys(args.embeddings)

    if args.queries_csv:
        queries = load_queries_from_csv(args.queries_csv, selected_values)
    else:
        queries = load_queries_from_xlsx(args.queries_xlsx, args.sheet_name, selected_values)
    queries = select_queries(queries, args.query_limit, args.one_per_dispute_type)

    if not queries:
        raise ValueError("검색할 질문이 없습니다. 입력 파일과 선택 값을 확인해주세요.")

    output_dir = make_output_dir(args.output_dir)
    all_candidate_rows: list[dict[str, Any]] = []
    all_chunk_rows: list[dict[str, Any]] = []
    manifests: dict[str, Any] = {}

    print(f"후보 풀링 시작: 질문 {len(queries)}개, top_k={args.top_k}")
    print(f"사용 임베딩: {', '.join(embedding_keys)}")

    collections: dict[str, Any] = {}
    embedders: dict[str, tuple[str, Any, str]] = {}
    for embedding_key in embedding_keys:
        collection, manifest = load_chroma_collection(args.vector_db_root, embedding_key)
        manifests[embedding_key] = manifest
        provider = manifest["embedding_provider"]
        model_name = manifest["embedding_model"]
        if provider == "sentence_transformers":
            embedder = load_sentence_transformer(model_name)
        else:
            embedder = load_openai_client()
        collections[embedding_key] = collection
        embedders[embedding_key] = (provider, embedder, model_name)
        print(f"- {embedding_key}: {model_name} 로드 완료")

    start_time = time.time()
    query_stats: list[dict[str, Any]] = []
    for query_index, query in enumerate(queries, start=1):
        pooled: dict[str, CandidateCase] = {}
        query_start = time.time()

        for embedding_key in embedding_keys:
            provider, embedder, model_name = embedders[embedding_key]
            query_vector = embed_query(provider, embedder, model_name, query.query_candidate)
            search_result = search_collection(
                collections[embedding_key],
                query_vector,
                args.top_k,
            )
            build_candidate_pool(query, embedding_key, search_result, pooled)

        candidates = sort_candidates(pooled)
        candidate_rows, chunk_rows = build_output_rows(
            query,
            candidates,
            args.final_cases_dir,
            args.preview_chars,
        )
        all_candidate_rows.extend(candidate_rows)
        all_chunk_rows.extend(chunk_rows)

        elapsed = time.time() - query_start
        query_stats.append(
            {
                "query_id": query.query_id,
                "candidate_count": len(candidates),
                "matched_chunk_count": sum(len(candidate.chunks) for candidate in candidates),
                "elapsed_seconds": round(elapsed, 2),
            }
        )
        print(
            f"[{query_index}/{len(queries)}] {query.query_id} "
            f"후보판례 {len(candidates)}건, 검색청크 {query_stats[-1]['matched_chunk_count']}건, "
            f"{elapsed:.1f}초",
            flush=True,
        )

    candidate_fields = [
        "query_id",
        "dispute_type",
        "query_candidate",
        "query_difficulty",
        "scenario_tags",
        "candidate_rank",
        "precedent_id",
        "case_no",
        "case_name",
        "court",
        "decision_date",
        "generated_summary",
        "matched_models",
        "best_rank",
        "best_distance",
        "matched_chunk_count",
        "gold_label",
        "label_reason",
    ]
    chunk_fields = [
        "query_id",
        "precedent_id",
        "candidate_rank",
        "model",
        "rank",
        "chunk_id",
        "chunk_type",
        "distance",
        "chunk_text_preview",
    ]
    write_csv(output_dir / "candidate_precedents.csv", all_candidate_rows, candidate_fields)
    write_csv(output_dir / "matched_chunks.csv", all_chunk_rows, chunk_fields)

    manifest = {
        "schema_version": "precedent_candidate_pooling.v1",
        "created_at": now_utc_iso(),
        "queries_count": len(queries),
        "top_k": args.top_k,
        "embedding_keys": embedding_keys,
        "vector_db_root": str(args.vector_db_root),
        "final_cases_dir": str(args.final_cases_dir),
        "candidate_rows": len(all_candidate_rows),
        "chunk_rows": len(all_chunk_rows),
        "elapsed_seconds": round(time.time() - start_time, 2),
        "vector_db_manifests": manifests,
        "query_stats": query_stats,
    }
    (output_dir / "pooling_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    print(f"완료: {output_dir}")
    print(f"- 후보 판례 CSV: {output_dir / 'candidate_precedents.csv'}")
    print(f"- 검색 청크 CSV: {output_dir / 'matched_chunks.csv'}")


if __name__ == "__main__":
    main()
